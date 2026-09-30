from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

slides = [
('Slide 1 — Introduction', [
'Good morning everyone.',
'My name is Omar Hammoud, and this project is actually very personal to me.',
'Alongside my background in Computer Science, I am also a volunteer assistant paramedic with the Islamic Medical Association in Saida.',
'So the idea for NGO Hub did not come only from thinking about what application I could build for my final project. It came from something I actually experience.',
'As volunteers, we go on ambulance missions, use different vehicles and medical equipment, complete reports, and try to keep information organized while also focusing on the most important thing — helping people.',
'That is why I built NGO Hub.'
]),
('Slide 2 — The Idea', [
'NGO Hub is a platform designed to make the daily work of a humanitarian healthcare organization easier.',
'Instead of having information spread between messages, papers, different files, or different people, the organization can have one place for its work.',
'It can organize ambulance missions, vehicles, medical equipment, staff activity, and reports.',
'It works in both Arabic and English, because I wanted it to be practical for the people who would actually use it.',
'But I do not want to spend too much time today showing you forms and tables. I want to show you the part of the project that I believe can save volunteers and staff the most time.',
'The AI assistant.'
]),
('Slide 3 — Why AI?', [
'When we return from a mission, or when there is a problem with an ambulance, someone still has to document what happened.',
'Usually, that person already knows the information. The problem is turning that information into an organized report.',
'So instead of asking the volunteer to go through many fields one by one, I wanted them to simply explain what happened naturally. The AI helps turn that explanation into an organized draft.',
'Let me show you a simple example.'
]),
('Slide 4 — AI Example 1: Vehicle Problem', [
'Imagine I am on duty and I notice a problem with one of our ambulances.',
'Instead of opening a long form and filling everything manually, I can write something simple like:',
'“Ambulance A03 has broken rear lights. The ambulance is still working, but it should be checked before the next night shift.”',
'That is it.',
'From this short message, the assistant understands what I am talking about and prepares the report for me. It identifies the ambulance, what the problem is, and how urgent it is.',
'Then I can review everything before submitting it.',
'The important point is that the AI does not make the decision for us. It helps us with the paperwork. The volunteer or staff member is still the person in control.'
]),
('Slide 5 — AI Example 2: Missing Information', [
'The second example is even more important. AI should not invent information.',
'Imagine I write:',
'“We completed the mission and transported the patient safely.”',
'But I forgot to mention where the mission happened. The assistant recognizes that something important is missing. Instead of guessing, it asks me:',
'“Where did the mission take place?”',
'I answer the question, and then it completes the draft.',
'For me, this is an important part of using AI responsibly. The system should help the volunteer, not replace their judgment.',
'It does not diagnose patients. It does not make medical decisions. And it does not automatically save something that the person did not review.'
]),
('Slide 6 — The Humanitarian Side', [
'And this is really the main idea behind NGO Hub.',
'When people hear “AI,” sometimes the conversation immediately becomes about the technology. But for this project, the technology is not the goal. The people are the goal.',
'I work with volunteers who give their time to respond to emergencies and help their community. Some of them may finish a long shift or return from a difficult mission and still have administrative work to complete.',
'If technology can reduce even a small part of that work, that means more time and attention can go toward the humanitarian mission itself.',
'That is what I wanted to build. Not AI that replaces volunteers. AI that supports them.'
]),
('Slide 7 — Closing', [
'NGO Hub started from a problem that I personally saw as a volunteer. And through what I learned in the Digital Hub, I was able to turn that problem into a real application.',
'Today, it can organize missions, vehicles, equipment, staff activity and reports, while also using AI to make reporting easier.',
'And this project can continue growing. For example, in the future, people could use the public website to contact the organization, ask about services, or even apply to volunteer.',
'For me, the success of NGO Hub is not about how many features it has. It is much simpler than that.',
'If it saves a volunteer a few minutes of paperwork…',
'if it helps information stay organized…',
'and if that gives the team a little more time to focus on helping someone who needs them…',
'then the project has done what I wanted it to do.',
'Thank you.'
])]

doc=Document()
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
sec=doc.sections[0]
sec.page_width=Inches(8.27)
sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.7)
sec.left_margin=sec.right_margin=Inches(.85)
for name in ['Normal','Title','Subtitle','Heading 1']:
    style=doc.styles[name]
    style.font.name='Calibri'
    style.font.color.rgb=RGBColor(0,0,0)
normal=doc.styles['Normal']
normal.font.size=Pt(12)
normal.paragraph_format.line_spacing=1.08
normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.widow_control=True
heading=doc.styles['Heading 1']
heading.font.size=Pt(16)
heading.font.bold=True
heading.paragraph_format.space_before=Pt(14)
heading.paragraph_format.space_after=Pt(8)
doc.styles['Title'].font.size=Pt(25)
doc.add_paragraph('NGO Hub Presentation Script','Title')
p=doc.add_paragraph('Omar Hammoud','Subtitle')
p.paragraph_format.space_after=Pt(12)
for i,(title,paras) in enumerate(slides):
    h=doc.add_paragraph(title,'Heading 1')
    if i in [2,4,6]:
        h.paragraph_format.page_break_before=True
    for text in paras:
        p=doc.add_paragraph(text)
        if text.startswith('“'):
            p.paragraph_format.left_indent=Inches(.2)
            p.runs[0].italic=True
footer=sec.footer.paragraphs[0]
footer.alignment=2
r=footer.add_run()
r.font.size=Pt(9)
field=OxmlElement('w:fldSimple')
field.set(qn('w:instr'),'PAGE')
r._r.addnext(field)
doc.core_properties.title='NGO Hub Presentation Script'
doc.core_properties.author='Omar Hammoud'
out=Path('NGO Hub Presentation Script.docx').resolve()
doc.save(out)
print(out)
